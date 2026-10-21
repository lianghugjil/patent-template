# -*- coding: utf-8 -*-
"""fix_layout.py — 套版式之后的收口脚本（幂等，可反复跑）。

放在 `apply_patent_format.py` **之后**执行：

    python apply_patent_format.py 专利.docx
    python fix_layout.py 专利.docx                  # 默认：行距收紧 + 附图铺满版心 + 附图分页
    python fix_layout.py 专利.docx --no-paginate    # 不分页（图小、想省页数时用）

做三件事（都是 apply_patent_format.py 不做、但成品必需）：

1. **行距收紧**：含"纯行内公式"的正文段落，行距由 atLeast 改回 exact 23 pt。
   （含独立公式 oMathPara 或分式/根式/积分/上下标堆叠结构的段落保留 atLeast，否则公式会被裁切）
2. **附图铺满版心**：图片宽度统一放大到版心满宽，高度按原比例同步。
   版心宽度按 `第一节页宽 − 左 − 右` 自动算，**页边距一改这里自动跟着变**，不用记常数。
3. **附图分页**：第 2 张图起另起一页，一页一幅。
   理由：附图放大到版心后，两张同页会挤到看不清标号；《审查指南》允许一页多幅，但清晰度优先。

为什么必须放大：默认插图宽度往往只有版心的 70% 左右，5 个标号挤在 8.9 cm 里，
审查员看标号、代理人答复改图都费劲。铺满后同样一张图高度从 8.9 cm 变 12.4 cm，标号明显清楚。
"""
import argparse
import sys

from docx import Document
from docx.shared import Pt, Mm
from docx.enum.text import WD_LINE_SPACING
from docx.oxml.ns import qn

M = "{http://schemas.openxmlformats.org/officeDocument/2006/math}"
WP = "{http://schemas.openxmlformats.org/drawingml/2006/wordprocessingDrawing}"
A = "{http://schemas.openxmlformats.org/drawingml/2006/main}"
BLIP = "{http://schemas.openxmlformats.org/drawingml/2006/main}blip"

# 这些 OMML 结构会把行撑高，行距必须留 atLeast
TALL = ("f", "rad", "nary", "limLow", "limUpp", "m", "groupChr")
LINE_EXACT = Pt(23)


def para_has_tall_math(p):
    for om in p._element.findall(".//" + M + "oMath"):
        for tag in TALL:
            if om.findall(".//" + M + tag):
                return True
    return False


def fix_leading(doc):
    fixed, kept = 0, 0
    for p in doc.paragraphs:
        el = p._element
        has_math = bool(el.findall(".//" + M + "oMath"))
        has_disp = bool(el.findall(".//" + M + "oMathPara"))
        has_img = bool(el.findall(".//" + qn("w:drawing")) or el.findall(".//" + qn("w:pict")))
        if not has_math or has_img:
            continue
        if has_disp or para_has_tall_math(p):
            kept += 1
            continue
        pf = p.paragraph_format
        if pf.line_spacing_rule == WD_LINE_SPACING.EXACTLY and pf.line_spacing == LINE_EXACT:
            continue
        pf.line_spacing = LINE_EXACT
        pf.line_spacing_rule = WD_LINE_SPACING.EXACTLY
        fixed += 1
    return fixed, kept


def fix_figure_width(doc, target_mm=None):
    """图片宽度铺满版心，高度等比。wp:extent 与 a:ext 必须同步，否则 Word 里变形。"""
    if target_mm is None:
        sec = doc.sections[0]
        target_mm = (sec.page_width - sec.left_margin - sec.right_margin) / 36000.0
    body = doc.element.body
    n = 0
    for ext in body.findall(".//" + WP + "extent"):
        cx, cy = int(ext.get("cx")), int(ext.get("cy"))
        if cx <= 0 or cy <= 0:
            continue
        new_cx = int(Mm(target_mm))
        new_cy = int(round(cy * new_cx / cx))
        ext.set("cx", str(new_cx))
        ext.set("cy", str(new_cy))
        inline = ext.getparent()
        for aext in inline.findall(".//" + A + "ext"):
            aext.set("cx", str(new_cx))
            aext.set("cy", str(new_cy))
        n += 1
    return n, target_mm


def paginate_figures(doc):
    """第 2 张图起另起一页。"""
    n = 0
    for p in doc.paragraphs:
        if p._element.findall(".//" + BLIP):
            n += 1
            if n >= 2:
                p.paragraph_format.page_break_before = True
    return n


def main():
    ap = argparse.ArgumentParser(add_help=True)
    ap.add_argument("path")
    ap.add_argument("--no-paginate", action="store_true", help="不做附图分页")
    ap.add_argument("--width", type=float, default=None, help="附图目标宽度 mm（默认铺满版心）")
    args = ap.parse_args()

    doc = Document(args.path)
    fixed, kept = fix_leading(doc)
    nfig, wmm = fix_figure_width(doc, args.width)
    n = 0 if args.no_paginate else paginate_figures(doc)
    doc.save(args.path)

    print(f"行距收紧: {fixed} 段 → exact 23 pt；保留 atLeast: {kept} 段（含分式/根式等）")
    print(f"附图放大: {nfig} 张 → 版心 {wmm:.1f} mm 宽（等比缩放）")
    if args.no_paginate:
        print("附图分页: 已跳过")
    else:
        print(f"附图分页: 共 {n} 张，第 2 张起各占一页")
    return 0


if __name__ == "__main__":
    sys.exit(main())
