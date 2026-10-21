#!/usr/bin/env python
# -*- coding: utf-8 -*-
"""apply_patent_format.py — 把「专利模板」版式套到任意 docx 上（页面/页眉/页脚/字体/行距/缩进）。

用法：
    python apply_patent_format.py 目标.docx                # 套版式（自动按册名切节）
    python apply_patent_format.py 目标.docx --report       # 只体检，不改文件
    python apply_patent_format.py 目标.docx --no-split     # 不切节，单节统一页眉

套用内容（数值全部来自已提交稿实测，见 references/docx_format.md）：
  页面 A4 / 上2.5 下2.0 左2.5 右2.0 cm / 页眉距1.5cm / 页脚距1.75cm / 行网格312
  页眉 黑体16pt 居中，文字为册名（说　明　书 等），字间全角空格
  页脚 宋体12pt 居中 PAGE 域，每册页码从 1 起
  正文 宋体12pt/Times New Roman，行距固定23pt，首行缩进2字符，两端对齐
  章节标题 宋体12pt 加粗、无缩进
  发明名称 居中不加粗，且用 w:ind 左-142/右+142 抵消左右边距不对称，落在纸张正中
  图片/独立公式 段落 最小值23pt（避免裁切）
"""
import sys
import os
import re
import argparse
from copy import deepcopy

from docx import Document
from docx.shared import Pt
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_LINE_SPACING
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.text.paragraph import Paragraph

# ---------------------------------------------------------------- 常量
VOLUMES = ('权利要求书', '说明书摘要', '说明书附图', '说明书', '摘要附图')
HEADINGS = ('技术领域', '背景技术', '发明内容', '附图说明', '具体实施方式', '实施例')

PAGE_W, PAGE_H = 11906, 16838          # twips, A4
# 页边距：上 2.5 / 下 2 / 左 2.5 / 右 2 cm（2026-10-21 定稿）
M_TOP, M_BOT, M_LEFT, M_RIGHT = 1418, 1134, 1418, 1134
HDR_DIST, FTR_DIST = 851, 992
LINE_EXACT = Pt(23)                     # 固定行距 23pt
INDENT_2CH = Pt(24)                     # 2 字符（12pt × 2）
CN_FONT, EN_FONT, HDR_FONT = '宋体', 'Times New Roman', '黑体'

SECTPR_ORDER = ['headerReference', 'footerReference', 'footnotePr', 'endnotePr', 'type',
                'pgSz', 'pgMar', 'paperSrc', 'pgBorders', 'lnNumType', 'pgNumType', 'cols',
                'formProt', 'vAlign', 'noEndnote', 'titlePg', 'textDirection', 'bidi',
                'rtlGutter', 'docGrid', 'printerSettings', 'sectPrChange']


def hdr_text(name):
    """册名 -> 页眉文字（字间全角空格）"""
    return '　'.join(list(name))


# ---------------------------------------------------------------- 小工具
def set_font(run, cn=CN_FONT, size=12, en=EN_FONT, bold=None):
    run.font.size = Pt(size)
    if bold is not None:
        run.font.bold = bold
    rPr = run._element.get_or_add_rPr()
    rFonts = rPr.find(qn('w:rFonts'))
    if rFonts is None:
        rFonts = OxmlElement('w:rFonts')
        rPr.insert(0, rFonts)
    rFonts.set(qn('w:ascii'), en)
    rFonts.set(qn('w:hAnsi'), en)
    rFonts.set(qn('w:eastAsia'), cn)
    rFonts.set(qn('w:cs'), en)


def set_sectpr_child(sectPr, tag, attrs):
    el = sectPr.find(qn('w:' + tag))
    if el is None:
        el = OxmlElement('w:' + tag)
        idx = SECTPR_ORDER.index(tag)
        inserted = False
        for child in list(sectPr):
            ctag = child.tag.split('}')[-1]
            if ctag in SECTPR_ORDER and SECTPR_ORDER.index(ctag) > idx:
                child.addprevious(el)
                inserted = True
                break
        if not inserted:
            sectPr.append(el)
    for k, v in attrs.items():
        el.set(qn('w:' + k), v)
    return el


def para_setup(p, align=None, indent=None, spacing='exact', before=0, after=0, bold=None):
    pf = p.paragraph_format
    if align is not None:
        pf.alignment = align
    pf.space_before = Pt(before)
    pf.space_after = Pt(after)
    if spacing == 'exact':
        pf.line_spacing = LINE_EXACT
        pf.line_spacing_rule = WD_LINE_SPACING.EXACTLY
    elif spacing == 'atleast':
        pf.line_spacing = LINE_EXACT
        pf.line_spacing_rule = WD_LINE_SPACING.AT_LEAST
    elif spacing == 'single':
        pf.line_spacing = 1.0
        pf.line_spacing_rule = WD_LINE_SPACING.SINGLE
    # 首行缩进（字符 + 绝对值双写，Word 按字符、其他渲染器按绝对值）
    if indent is None:
        pf.first_line_indent = None
        pPr = p._element.find(qn('w:pPr'))
        if pPr is not None:
            ind = pPr.find(qn('w:ind'))
            if ind is not None:
                pPr.remove(ind)
    else:
        pf.first_line_indent = indent
        pPr = p._element.get_or_add_pPr()
        indEl = pPr.find(qn('w:ind'))
        chars = str(int(round(indent.pt / 12 * 100)))   # 12pt 字号下 1 字符 = 100
        indEl.set(qn('w:firstLineChars'), chars)
        for k in ('w:left', 'w:right', 'w:leftChars', 'w:rightChars'):
            indEl.attrib.pop(qn(k), None)
    if bold is not None:
        for r in p.runs:
            r.font.bold = bold


def center_on_page(p, sec):
    """把居中段落的光学中心对到纸张中心。

    Word 的“居中”是在**版心**内居中；页边距左 2.5 / 右 2 cm 不对称时，版心中心比纸张
    中心偏右 0.25 cm，标题看上去就“没居中”。这里给标题段补一对负左缩进 / 正右缩进，
    把居中窗口整体左移 (右-左)/2，使标题落在纸张正中。
    """
    off = int(round((sec.right_margin.twips - sec.left_margin.twips) / 2))
    pPr = p._element.get_or_add_pPr()
    ind = pPr.find(qn('w:ind'))
    if ind is None:
        ind = OxmlElement('w:ind')
        _insert_ppr_child(pPr, ind, 'ind')
    ind.set(qn('w:left'), str(off))
    ind.set(qn('w:right'), str(-off))
    ind.set(qn('w:firstLine'), '0')
    ind.set(qn('w:firstLineChars'), '0')
    return off


PPR_ORDER = ['pStyle', 'keepNext', 'keepLines', 'pageBreakBefore', 'framePr', 'widowControl',
             'numPr', 'suppressLineNumbers', 'pBdr', 'shd', 'tabs', 'suppressAutoHyphens',
             'kinsoku', 'wordWrap', 'overflowPunct', 'topLinePunct', 'autoSpaceDE',
             'autoSpaceDN', 'bidi', 'adjustRightInd', 'snapToGrid', 'spacing', 'ind',
             'contextualSpacing', 'mirrorIndents', 'suppressOverlap', 'jc', 'textDirection',
             'textAlignment', 'textboxTightWrap', 'outlineLvl', 'divId', 'cnfStyle', 'rPr',
             'sectPr', 'pPrChange']


def _insert_ppr_child(pPr, el, tag):
    """按 OOXML 顺序插入 pPr 子元素（否则 Word/LibreOffice 可能忽略该属性）。"""
    idx = PPR_ORDER.index(tag)
    for child in list(pPr):
        ctag = child.tag.split('}')[-1]
        if ctag in PPR_ORDER and PPR_ORDER.index(ctag) > idx:
            child.addprevious(el)
            return
    pPr.append(el)


def own_header_names(doc):
    """各节**自己**的页眉册名（不走 linkToPrevious 继承），用于幂等重跑时还原顺序。"""
    out = []
    for sec in doc.sections:
        txt = ''
        sectPr = sec._sectPr
        for ref in sectPr.findall(qn('w:headerReference')):
            if ref.get(qn('w:type')) not in (None, 'default'):
                continue
            rid = ref.get(qn('r:id'))
            part = sec.part.related_parts.get(rid)
            if part is not None:
                xml = part.blob.decode('utf-8', 'ignore')
                txt = ''.join(re.findall(r'<w:t[^>]*>([^<]*)</w:t>', xml))
            break
        out.append(txt.replace('　', '').strip())
    return out


def locate_title(doc, names):
    """定位说明书首页的发明名称段：说明书册的第一段非空正文（≤40 字、不以句号结尾）。"""
    if '说明书' not in names:
        return None
    target = names.index('说明书')
    body = doc.element.body
    seen = 0                     # 已跨过的分节符数
    for el in body.iterchildren():
        if el.tag != qn('w:p'):
            continue
        pPr = el.find(qn('w:pPr'))
        sect = pPr.find(qn('w:sectPr')) if pPr is not None else None
        txt = ''.join(t.text or '' for t in el.iter(qn('w:t'))).strip()
        if sect is not None:     # 这一节到此结束
            seen += 1
            continue
        if seen != target or not txt:
            continue
        if txt in HEADINGS or txt.startswith('图'):
            return None
        if re.match(r'^\d{1,2}\s*[.．、]', txt):     # 权利要求编号段，不是标题
            return None
        if len(txt) <= 40 and not txt.endswith(('。', '；', '，')):
            return el
        return None
    return None


def add_page_field(p):
    def mk(tag=None, text=None):
        r = p.add_run()
        if tag:
            el = OxmlElement('w:fldChar')
            el.set(qn('w:fldCharType'), tag)
            r._element.append(el)
        if text is not None:
            el = OxmlElement('w:instrText')
            el.set(qn('xml:space'), 'preserve')
            el.text = text
            r._element.append(el)
        set_font(r, CN_FONT, 12)
        return r
    mk('begin')
    mk(text=' PAGE ')
    mk('separate')
    r = p.add_run('1')
    set_font(r, CN_FONT, 12)
    mk('end')


def clear_para(p):
    for r in list(p.runs):
        r._element.getparent().remove(r._element)
    for child in list(p._element):
        if child.tag.split('}')[-1] not in ('pPr',):
            p._element.remove(child)


def is_empty_p(el):
    if el.findall('.//' + qn('w:drawing')) or el.findall('.//' + qn('w:pict')):
        return False
    if el.findall('.//' + qn('m:oMath')):
        return False
    return not ''.join(t.text or '' for t in el.findall('.//' + qn('w:t'))).strip()


# ---------------------------------------------------------------- 切节
def split_sections(doc):
    """按册名段切节：册名文字移入页眉，正文里删掉册名行。返回册名顺序列表。"""
    body = doc.element.body
    final_sectPr = body.find(qn('w:sectPr'))
    paras = [p for p in doc.paragraphs]
    names = []
    first_content_seen = False
    for p in paras:
        txt = p.text.strip()
        el = p._element
        if txt in VOLUMES:
            names.append(txt)
            if not first_content_seen:
                # 文档开头的册名：直接删行，本节页眉用该册名
                el.getparent().remove(el)
                continue
            # 先清掉前面紧邻的空段，避免多出空白
            prev = el.getprevious()
            while prev is not None and prev.tag == qn('w:p') and is_empty_p(prev):
                nxt = prev.getprevious()
                prev.getparent().remove(prev)
                prev = nxt
            # 插入带 sectPr 的空段收束上一节
            carrier = OxmlElement('w:p')
            pPr = OxmlElement('w:pPr')
            spacing = OxmlElement('w:spacing')
            spacing.set(qn('w:before'), '0')
            spacing.set(qn('w:after'), '0')
            spacing.set(qn('w:line'), '20')
            spacing.set(qn('w:lineRule'), 'exact')
            pPr.append(spacing)
            rPr = OxmlElement('w:rPr')
            sz = OxmlElement('w:sz')
            sz.set(qn('w:val'), '2')
            rPr.append(sz)
            pPr.append(rPr)
            sp = deepcopy(final_sectPr)
            for item in (sp.findall(qn('w:headerReference')) + sp.findall(qn('w:footerReference'))):
                sp.remove(item)
            pPr.append(sp)
            carrier.append(pPr)
            el.addprevious(carrier)
            el.getparent().remove(el)
            first_content_seen = True
        elif txt:
            first_content_seen = True
    return names


# ---------------------------------------------------------------- 页眉页脚
#: 页眉下边框 —— 与历史提交稿完全一致：single、0.75 pt（sz 单位 1/8 pt）、距文字 1 pt
HDR_BORDER = {'val': 'single', 'sz': '6', 'space': '1', 'color': 'auto'}


def header_border(p):
    """页眉段落写死**只有下边框**、上左右及其余边置 nil。

    Word 内置 Header 样式与历史稿的 header 样式都自带
    ``<w:bottom w:val="single" w:sz="6" w:space="1"/>``（专利稿页眉下面那道横线）。
    段落级显式写一遍，既不依赖样式是否定义，也保证在别的 Word/WPS 里一致。
    """
    pPr = p._element.get_or_add_pPr()
    for old in pPr.findall(qn('w:pBdr')):
        pPr.remove(old)
    pBdr = OxmlElement('w:pBdr')
    for edge in ('top', 'left', 'bottom', 'right', 'between', 'bar'):
        e = OxmlElement('w:' + edge)
        if edge == 'bottom':
            for k, v in HDR_BORDER.items():
                e.set(qn('w:' + k), v)
        else:
            e.set(qn('w:val'), 'nil')
        pBdr.append(e)
    ps = pPr.find(qn('w:pStyle'))
    if ps is not None:
        ps.addnext(pBdr)
    else:
        pPr.insert(0, pBdr)


def no_border(p):
    """清掉段落边框（页脚用，页码下面不要横线）。"""
    pPr = p._element.get_or_add_pPr()
    for old in pPr.findall(qn('w:pBdr')):
        pPr.remove(old)
    pBdr = OxmlElement('w:pBdr')
    for edge in ('top', 'left', 'bottom', 'right', 'between', 'bar'):
        e = OxmlElement('w:' + edge)
        e.set(qn('w:val'), 'nil')
        e.set(qn('w:sz'), '0')
        e.set(qn('w:space'), '0')
        pBdr.append(e)
    ps = pPr.find(qn('w:pStyle'))
    if ps is not None:
        ps.addnext(pBdr)
    else:
        pPr.insert(0, pBdr)


def has_border(p, edge=None):
    """段落是否有可见边框；指定 edge 时只看该边是否为 single。"""
    pBdr = p._element.find(qn('w:pPr') + '/' + qn('w:pBdr'))
    if pBdr is None:
        return False
    hits = []
    for e in pBdr:
        val = e.get(qn('w:val'))
        if val in (None, 'nil', 'none'):
            continue
        tag = e.tag.split('}')[-1]
        if edge is None or tag == edge:
            hits.append(tag)
    return bool(hits)


def set_header(sec, name):
    sec.header.is_linked_to_previous = False
    hdr = sec.header
    for p in list(hdr.paragraphs)[1:]:
        p._element.getparent().remove(p._element)
    p = hdr.paragraphs[0]
    clear_para(p)
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    pf = p.paragraph_format
    pf.space_before = Pt(0)
    pf.space_after = Pt(0)
    pf.line_spacing = 1.0
    pf.line_spacing_rule = WD_LINE_SPACING.SINGLE
    pf.first_line_indent = None
    run = p.add_run(hdr_text(name))
    set_font(run, HDR_FONT, 16, en=HDR_FONT)
    header_border(p)
    center_on_page(p, sec)          # 册名与专利标题同轴（纸张中心）


def set_footer(sec):
    sec.footer.is_linked_to_previous = False
    ftr = sec.footer
    for p in list(ftr.paragraphs)[1:]:
        p._element.getparent().remove(p._element)
    p = ftr.paragraphs[0]
    clear_para(p)
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    pf = p.paragraph_format
    pf.space_before = Pt(0)
    pf.space_after = Pt(0)
    pf.line_spacing = 1.0
    pf.line_spacing_rule = WD_LINE_SPACING.SINGLE
    pf.first_line_indent = None
    add_page_field(p)
    center_on_page(p, sec)          # 页码与页眉同轴
    no_border(p)


def setup_section(sec):
    sec.page_width = PAGE_W * 635          # twips -> EMU (1 twip = 635 EMU)
    sec.page_height = PAGE_H * 635
    sec.top_margin = M_TOP * 635
    sec.bottom_margin = M_BOT * 635
    sec.left_margin = M_LEFT * 635
    sec.right_margin = M_RIGHT * 635
    sec.header_distance = HDR_DIST * 635
    sec.footer_distance = FTR_DIST * 635
    sectPr = sec._sectPr
    set_sectpr_child(sectPr, 'docGrid', {'type': 'lines', 'linePitch': '312'})
    set_sectpr_child(sectPr, 'pgNumType', {'start': '1'})


# ---------------------------------------------------------------- 正文
def style_normal(doc):
    st = doc.styles['Normal']
    st.font.name = EN_FONT
    st.font.size = Pt(12)
    rPr = st.element.get_or_add_rPr()
    rFonts = rPr.find(qn('w:rFonts'))
    if rFonts is None:
        rFonts = OxmlElement('w:rFonts')
        rPr.insert(0, rFonts)
    rFonts.set(qn('w:ascii'), EN_FONT)
    rFonts.set(qn('w:hAnsi'), EN_FONT)
    rFonts.set(qn('w:eastAsia'), CN_FONT)
    rFonts.set(qn('w:cs'), EN_FONT)
    st.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    # 关闭孤行控制（与语料一致）
    pPr = st.element.get_or_add_pPr()
    wc = pPr.find(qn('w:widowControl'))
    if wc is None:
        wc = OxmlElement('w:widowControl')
        pPr.insert(0, wc)
    wc.set(qn('w:val'), '0')


def prune_empty_paragraphs(doc):
    """删掉正文空段（模板段落之间不空行），只在章节标题前保留一个整行空行。"""
    paras = [p for p in doc.paragraphs]
    removed = 0
    i = 0
    while i < len(paras):
        el = paras[i]._element
        if is_empty_p(el) and el.find('.//' + qn('w:sectPr')) is None:
            run = []
            j = i
            while j < len(paras):
                e = paras[j]._element
                if is_empty_p(e) and e.find('.//' + qn('w:sectPr')) is None:
                    run.append(paras[j])
                    j += 1
                else:
                    break
            nxt = paras[j] if j < len(paras) else None
            keep = nxt is not None and nxt.text.strip() in HEADINGS
            for k, q in enumerate(run):
                if keep and k == len(run) - 1:
                    q.paragraph_format.space_before = Pt(0)
                    q.paragraph_format.space_after = Pt(0)
                    q.paragraph_format.line_spacing = LINE_EXACT
                    q.paragraph_format.line_spacing_rule = WD_LINE_SPACING.EXACTLY
                    q.paragraph_format.first_line_indent = None
                else:
                    q._element.getparent().remove(q._element)
                    removed += 1
            i = j
        else:
            i += 1
    return removed


def format_body(doc, title_el=None, title_sec=None):
    pending_title = False
    stats = {'body': 0, 'heading': 0, 'title': 0, 'img': 0, 'math': 0, 'blank': 0}
    stats['blank'] = prune_empty_paragraphs(doc)
    for p in doc.paragraphs:
        el = p._element
        txt = p.text.strip()
        has_img = bool(el.findall('.//' + qn('w:drawing')) or el.findall('.//' + qn('w:pict')))
        has_disp = bool(el.findall('.//' + qn('m:oMathPara')))
        has_math = bool(el.findall('.//' + qn('m:oMath')))

        if txt == '发明名称':
            el.getparent().remove(el)
            pending_title = True
            continue
        if not txt and not has_img and not has_math:
            continue

        if pending_title or (title_el is not None and el is title_el):
            para_setup(p, WD_ALIGN_PARAGRAPH.CENTER, None, 'exact')
            for r in p.runs:
                set_font(r, CN_FONT, 12, bold=False)
            if title_sec is not None:
                center_on_page(p, title_sec)
            pending_title = False
            stats['title'] += 1
            continue

        if txt in HEADINGS:
            para_setup(p, WD_ALIGN_PARAGRAPH.JUSTIFY, Pt(0), 'exact')
            for r in p.runs:
                set_font(r, CN_FONT, 12, bold=True)
            stats['heading'] += 1
            continue

        if has_img:
            para_setup(p, WD_ALIGN_PARAGRAPH.CENTER, None, 'atleast', before=6, after=6)
            stats['img'] += 1
            continue

        if (has_disp or has_math) and not txt:
            para_setup(p, WD_ALIGN_PARAGRAPH.CENTER, None, 'atleast', before=6, after=6)
            stats['math'] += 1
            continue

        if txt and len(txt) <= 6 and txt.startswith('图') and txt[1:].isdigit():
            para_setup(p, WD_ALIGN_PARAGRAPH.CENTER, None, 'exact')
            for r in p.runs:
                set_font(r, CN_FONT, 12)
            continue

        para_setup(p, WD_ALIGN_PARAGRAPH.JUSTIFY, INDENT_2CH,
                   'atleast' if has_math else 'exact')
        for r in p.runs:
            set_font(r, CN_FONT, 12)
        stats['body'] += 1
    return stats


def format_tables(doc):
    n = 0
    for t in doc.tables:
        for ri, row in enumerate(t.rows):
            for cell in row.cells:
                for p in cell.paragraphs:
                    para_setup(p, WD_ALIGN_PARAGRAPH.CENTER, None, 'single')
                    for r in p.runs:
                        set_font(r, CN_FONT, 10.5, bold=(True if ri == 0 else None))
        n += 1
    return n


# ---------------------------------------------------------------- 主流程
def apply_format(path, split=True, volume=None):
    doc = Document(path)
    names = split_sections(doc) if split else []
    if not names:
        existing = own_header_names(doc)
        if any(existing):                    # 已切过节：按现有页眉册名还原顺序（幂等）
            names = [n or (volume or '说明书') for n in existing]
        else:
            names = [volume or '说明书']

    style_normal(doc)
    for i, sec in enumerate(doc.sections):
        setup_section(sec)
        name = names[i] if i < len(names) else names[-1]
        set_header(sec, name)
        set_footer(sec)

    # 说明书首页的发明名称：定位到具体段落，居中并对齐到纸张中心
    title_el, title_sec = None, None
    if '说明书' in names:
        title_el = locate_title(doc, names)
        if title_el is not None:
            title_sec = doc.sections[names.index('说明书')]

    stats = format_body(doc, title_el, title_sec)
    nt = format_tables(doc)
    doc.save(path)
    return names, stats, nt


def report(path):
    doc = Document(path)
    ok = True
    print('=' * 66)
    print('文档:', os.path.basename(path))
    print('=' * 66)
    for i, sec in enumerate(doc.sections):
        h = [p.text for p in sec.header.paragraphs if p.text.strip()]
        f = [p.text for p in sec.footer.paragraphs if p.text.strip()]
        r = sec.header.paragraphs[0].runs if h else []
        hfont = (r[0].font.name, r[0].font.size.pt if r[0].font.size else None) if r else (None, None)
        print(f'第{i+1}节 页边距{t_m(sec.top_margin)}/{t_m(sec.bottom_margin)}/{t_m(sec.left_margin)}/{t_m(sec.right_margin)}')
        print(f'      页眉={h[0] if h else "(空)"!r}  字体={hfont}')
        print(f'      页脚={f[0] if f else "(空)"!r}')
        if (sec.top_margin.twips, sec.bottom_margin.twips, sec.left_margin.twips,
                sec.right_margin.twips) != (M_TOP, M_BOT, M_LEFT, M_RIGHT):
            ok = False
            print('      !! 页边距不符')
        if not h:
            ok = False
            print('      !! 缺页眉')
        if not f:
            ok = False
            print('      !! 缺页码')
        hp = sec.header.paragraphs[0] if sec.header.paragraphs else None
        fp = sec.footer.paragraphs[0] if sec.footer.paragraphs else None
        hb = has_border(hp, 'bottom') if hp is not None else False
        fb = has_border(fp) if fp is not None else False
        print(f'      边框=页眉下横线 {"有(0.75pt)" if hb else "无"}  页脚 {"无边框" if not fb else "有多余边框"}')
        if hp is not None and not hb:
            ok = False
            print('      !! 页眉缺下边框（专利稿页眉应有 0.75pt 横线）')
        if fp is not None and fb:
            ok = False
            print('      !! 页脚有多余边框（页码不应带横线）')

    # ---- 说明书首页发明名称居中 ----
    hdr_names = own_header_names(doc)
    if not any(hdr_names):
        hdr_names = [''.join(p.text for p in sec.header.paragraphs).replace('　', '').strip()
                     for sec in doc.sections]
    tel = locate_title(doc, hdr_names)
    if tel is not None:
        tp = Paragraph(tel, doc)
        c = is_center(tp)
        off = 0
        ind = tel.find(qn('w:pPr'))
        ind = ind.find(qn('w:ind')) if ind is not None else None
        if ind is not None:
            off = int(ind.get(qn('w:left')) or 0)
        print(f'发明名称  「{tp.text[:34]}」  居中={"是" if c else "否"}  纸张中心补偿={off} twips')
        if not c:
            ok = False
            print('      !! 说明书首页发明名称未居中')
    else:
        print('发明名称  未定位到说明书标题段（跳过）')

    n_bad_ind = n_bad_ls = n_bad_font = n_body = 0
    for p in doc.paragraphs:
        txt = p.text.strip()
        el = p._element
        if not txt or txt in HEADINGS:
            continue
        if el.findall('.//' + qn('w:drawing')):
            continue
        if is_center(p):
            continue
        n_body += 1
        pf = p.paragraph_format
        if pf.first_line_indent is None or abs(pf.first_line_indent.pt - 24) > 0.6:
            n_bad_ind += 1
        if pf.line_spacing is None or abs(float(pf.line_spacing) - LINE_EXACT) > 200:
            n_bad_ls += 1
        if p.runs:
            r = p.runs[0]
            if r.font.size is None or abs(r.font.size.pt - 12) > 0.1:
                n_bad_font += 1
    print(f'\n正文段 {n_body} 段：缩进异常 {n_bad_ind}，行距异常 {n_bad_ls}，字号异常 {n_bad_font}')
    if n_bad_ind or n_bad_ls or n_bad_font:
        ok = False
    print('\n结论:', '✅ 全部通过' if ok else '⚠️ 存在不符项')
    return ok


def is_center(p):
    al = p.paragraph_format.alignment
    return al is not None and 'CENTER' in str(al)


def t_m(v):
    return f'{v.cm:.1f}cm' if v is not None else '?'


if __name__ == '__main__':
    ap = argparse.ArgumentParser()
    ap.add_argument('docx')
    ap.add_argument('--report', action='store_true')
    ap.add_argument('--no-split', action='store_true')
    ap.add_argument('--volume', default=None)
    a = ap.parse_args()
    if a.report:
        report(a.docx)
    else:
        names, stats, nt = apply_format(a.docx, split=not a.no_split, volume=a.volume)
        print('册名 -> 页眉:', ' | '.join(hdr_text(n) for n in names))
        print('段落统计:', stats, '| 表格:', nt)
        report(a.docx)
