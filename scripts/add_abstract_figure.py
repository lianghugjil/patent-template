# -*- coding: utf-8 -*-
"""add_abstract_figure.py —— 给已套版式的申请文件补「摘要附图」册。

    python add_abstract_figure.py 目标.docx                        # 取说明书附图册第一幅，宽 10 cm
    python add_abstract_figure.py 目标.docx --width-cm 9.5
    python add_abstract_figure.py 目标.docx --move-abstract-front   # 顺手把「说明书摘要」挪到卷首
    python add_abstract_figure.py 目标.docx --image 附图1.png        # 指定用哪一幅

为什么必须有这一册
------------------
审查指南第一部分第一章 4.5.2：有附图的申请须提交一幅最能说明技术方案主要技术特征的附图作为
摘要附图，**且摘要附图应当是说明书附图中的一幅**；细则 26 条：缩到 4 cm × 6 cm 仍能分辨细节。
官方装订顺序（《专利申请受理和审批办事指南》）：
    请求书 → **说明书摘要 → 摘要附图** → 权利要求书 → 说明书 → 说明书附图

本脚本干三件事
--------------
1. 册序：若「说明书摘要」排在「权利要求书」之后（早期模板的历史遗留），按官方顺序挪到卷首；
2. 补册：在「说明书摘要」之后插入一节，页眉 `摘　要　附　图`、单独编页、只放一幅图，
   图幅按版心宽度 60%~70%（默认 10 cm）**居中、不铺满**，页内不写任何文字（含图注）；
3. 幂等：已经存在「摘要附图」册时只替换其中的图片，不重复插节。

图从哪来：默认直接抠**本文档「说明书附图」册里的第一幅**（即附图 1）——这样保证"摘要附图必须是
说明书附图中的一幅"，也不会因为外部 png 改名/丢失而失效。`--image` 可显式指定外部图片。
"""
from __future__ import annotations

import argparse
import copy
import io
import os
import sys

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.shared import Cm, Pt

sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))
from apply_patent_format import (set_footer, set_header, set_sectpr_child,   # noqa: E402
                                setup_section)

ABS = '说明书摘要'
AFG = '摘要附图'
CLM = '权利要求书'
DWG = '说明书附图'

MAX_H_CM = 20.0         # 图高上限（版心高 25.2 cm，上下留呼吸）


# ---------------------------------------------------------------- 分块
def hdr_name(sec):
    return ''.join(p.text for p in sec.header.paragraphs).replace('\u3000', '').strip()


def blocks(doc):
    """按「节」把 body 子元素切块。返回 [(册名, [元素...]), ...]，元素顺序即正文顺序。"""
    body = doc.element.body
    kids = [k for k in body if k.tag != qn('w:sectPr')]
    carriers = []
    for i, k in enumerate(kids):
        if k.tag != qn('w:p'):
            continue
        pPr = k.find(qn('w:pPr'))
        if pPr is not None and pPr.find(qn('w:sectPr')) is not None:
            carriers.append(i)
    names = [hdr_name(s) for s in doc.sections]
    out, start = [], 0
    for n, c in enumerate(carriers):
        out.append((names[n] if n < len(names) else '', kids[start:c + 1]))
        start = c + 1
    if start < len(kids):
        out.append((names[-1] if names else '', kids[start:]))
    return out


def block_index(blks, name):
    for i, (n, _) in enumerate(blks):
        if n == name:
            return i
    return -1


# ---------------------------------------------------------------- 图
def first_image_blob(doc, blks):
    """抠「说明书附图」册里第一幅图的原图字节（没这一册就全文找）。"""
    order = [i for i, (n, _) in enumerate(blks) if n == DWG] + \
            [i for i in range(len(blks)) if blks[i][0] != DWG]
    for i in order:
        for el in blks[i][1]:
            blip = el.find('.//' + qn('a:blip'))
            if blip is None:
                continue
            rId = blip.get(qn('r:embed')) or blip.get(qn('r:link'))
            if rId:
                return doc.part.related_parts[rId].blob
    return None


def fit_size(blob, width_cm):
    """按宽度定尺寸，图太高则按高度上限反算宽度。返回 (w_cm, h_cm)。"""
    try:
        from PIL import Image
        im = Image.open(io.BytesIO(blob))
        w_px, h_px = im.size
        h = width_cm * h_px / float(w_px)
        if h > MAX_H_CM:
            h, width_cm = MAX_H_CM, MAX_H_CM * w_px / float(h_px)
        return width_cm, h
    except Exception:
        return width_cm, width_cm * 0.75


def _paras_in(doc, elements, with_pic):
    ids = {id(e) for e in elements}
    out = []
    for p in doc.paragraphs:
        if id(p._p) not in ids:
            continue
        if with_pic and p._p.find('.//' + qn('a:blip')) is None:
            continue
        out.append(p)
    return out


def _style_absfig_section(sec):
    """摘要附图节：同上版式 + 页眉 摘　要　附　图 + 单独编页 + 页面内容垂直居中。"""
    setup_section(sec)
    set_header(sec, AFG)
    set_footer(sec)
    set_sectpr_child(sec._sectPr, 'vAlign', {'val': 'center'})


# ---------------------------------------------------------------- 主流程
def main():
    ap = argparse.ArgumentParser()
    ap.add_argument('docx')
    ap.add_argument('--width-cm', type=float, default=10.0,
                    help='图幅宽度（cm），默认 10（版心 16.5 cm 的 61%%）')
    ap.add_argument('--image', default=None, help='指定图片；默认取本文档说明书附图册第一幅')
    ap.add_argument('--move-abstract-front', action='store_true',
                    help='「说明书摘要」若排在权利要求书之后，按官方装订顺序挪到卷首')
    ap.add_argument('--no-reorder', action='store_true', help='只补册，不动册序')
    args = ap.parse_args()

    path = os.path.abspath(args.docx)
    if not os.path.isfile(path):
        sys.exit('找不到文件：%s' % path)
    doc = Document(path)

    blks = blocks(doc)
    print('当前册序：', ' → '.join(n or '(无册名)' for n, _ in blks) or '(无)')

    # ---------- 1. 册序 ----------
    i_abs, i_afg, i_clm = (block_index(blks, ABS), block_index(blks, AFG),
                           block_index(blks, CLM))
    if i_abs < 0:
        sys.exit('没找到「说明书摘要」册——先跑 apply_patent_format.py 把册名切进页眉')
    if args.move_abstract_front and not args.no_reorder and i_clm >= 0 and i_abs > i_clm:
        first = doc.element.body[0]
        for el in blks[i_abs][1]:
            first.addprevious(el)
        print('已把「说明书摘要」挪到卷首（官方装订顺序：摘要 → 摘要附图 → 权利要求书）')
        blks = blocks(doc)
        print('调整后册序：', ' → '.join(n or '(无册名)' for n, _ in blks))
        i_abs, i_afg, i_clm = (block_index(blks, ABS), block_index(blks, AFG),
                               block_index(blks, CLM))

    # ---------- 2. 取图 ----------
    if args.image:
        with open(args.image, 'rb') as f:
            blob = f.read()
    else:
        blob = first_image_blob(doc, blks)
        if blob is None:
            sys.exit('全文没图——无附图可以不设摘要附图；要设请用 --image 指定图片')
    w_cm, h_cm = fit_size(blob, args.width_cm)

    # ---------- 3. 插册 / 替换图 ----------
    if i_afg >= 0:
        cand = _paras_in(doc, blks[i_afg][1], with_pic=True)
        if cand:
            target = cand[0]
            for p in cand[1:]:                 # 只留一幅
                p._p.getparent().remove(p._p)
            for r in list(target._p.findall(qn('w:r'))):
                target._p.remove(r)
            verb = '图幅改为'
        else:
            # 空册：节还在、图被删掉了（WPS 里删旧图时极易连带删掉摘要附图）。
            # 在该节 sectPr 载体段之前补一个图片段即可，不必重建整节。
            afg_carrier = blks[i_afg][1][-1]
            if afg_carrier.find(qn('w:pPr')) is None or \
                    afg_carrier.find(qn('w:pPr')).find(qn('w:sectPr')) is None:
                sys.exit('「摘要附图」册的末尾元素不是 sectPr 载体段——文档结构异常')
            target = doc.add_paragraph()
            afg_carrier.addprevious(target._p)
            verb = '空册已补图，图幅'
        target.alignment = WD_ALIGN_PARAGRAPH.CENTER
        target.paragraph_format.first_line_indent = None
        target.add_run().add_picture(io.BytesIO(blob), width=Cm(w_cm))
        _style_absfig_section(doc.sections[i_afg])
        print('「摘要附图」册已存在——%s %.1f cm × %.1f cm' % (verb, w_cm, h_cm))
    else:
        abs_carrier = blks[i_abs][1][-1]       # 摘要节的 sectPr 载体段
        pPr = abs_carrier.find(qn('w:pPr'))
        sect_src = pPr.find(qn('w:sectPr')) if pPr is not None else None
        if sect_src is None:
            sys.exit('「说明书摘要」节末尾没有 sectPr——文档结构异常')
        sect_new = copy.deepcopy(sect_src)
        for tag in ('w:headerReference', 'w:footerReference'):
            for e in sect_new.findall(qn(tag)):
                if (e.get(qn('w:type')) or 'default') == 'default':
                    sect_new.remove(e)

        img_p = doc.add_paragraph()            # 先建在末尾，再挪到摘要之后
        carry_p = doc.add_paragraph()
        abs_carrier.addnext(carry_p._p)
        abs_carrier.addnext(img_p._p)          # → [摘要末段][图片段][sectPr 载体段]
        carry_pPr = carry_p._p.get_or_add_pPr()
        for old in carry_pPr.findall(qn('w:sectPr')):
            carry_pPr.remove(old)
        carry_pPr.append(sect_new)

        img_p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        img_p.paragraph_format.first_line_indent = None
        img_p.paragraph_format.space_before = Pt(0)
        img_p.paragraph_format.space_after = Pt(0)
        img_p.add_run().add_picture(io.BytesIO(blob), width=Cm(w_cm))

        new_idx = i_abs + 1                    # 新节紧跟摘要节
        if new_idx >= len(doc.sections):
            sys.exit('插册失败：节数没增加（sectPr 载体段没插进去）')
        _style_absfig_section(doc.sections[new_idx])   # 此时这一节还在继承摘要的页眉
        i_afg = block_index(blocks(doc), AFG)
        if i_afg != new_idx:
            sys.exit('插册失败：册名没落到第 %d 节（实际 %d）' % (new_idx, i_afg))
        print('已插入「摘要附图」册：页眉 摘　要　附　图，单独编页，'
              '图幅 %.1f cm × %.1f cm 居中' % (w_cm, h_cm))

    # ---------- 4. 存盘 ----------
    tmp = os.path.splitext(path)[0] + '.absfig.docx'
    doc.save(tmp)
    try:
        os.replace(tmp, path)
    except PermissionError:
        sys.exit('!! 成品被占用（Word / 预览面板锁着），结果留在：%s' % tmp)

    print('最终册序：', ' → '.join(n or '(无册名)' for n, _ in blocks(Document(path))))
    print('✔ 完成：%s' % path)


if __name__ == '__main__':
    main()
